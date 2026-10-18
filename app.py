
from pathlib import Path
import warnings

import numpy as np
import pandas as pd
from scipy import stats
from sklearn.model_selection import train_test_split
from sklearn.linear_model import LinearRegression
from sklearn.metrics import r2_score, mean_squared_error
import sympy as sp
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import seaborn as sns

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
OUTPUT_DIR = BASE_DIR / "output"
CHART_DIR = OUTPUT_DIR / "charts"

OUTPUT_DIR.mkdir(exist_ok=True)
CHART_DIR.mkdir(exist_ok=True)

INPUT_FILE = DATA_DIR / "sales_data.xlsx"
OUTPUT_FILE = OUTPUT_DIR / "Sales_Analysis_Dashboard.xlsx"

sns.set_theme(style="whitegrid")


if not INPUT_FILE.exists():
    raise FileNotFoundError(
        f"Excel file not found: {INPUT_FILE}\n"
        "Download a sales Excel dataset and save it as "
        "data/sales_data.xlsx"
    )

excel = pd.ExcelFile(INPUT_FILE)
print("Available Excel sheets:", excel.sheet_names)


df = None

for sheet in excel.sheet_names:
    temp = pd.read_excel(INPUT_FILE, sheet_name=sheet)
    temp.columns = [str(c).strip() for c in temp.columns]

    if any(c.lower() in ["sales", "revenue", "total sales"]
           for c in temp.columns):
        df = temp
        print("Reading sheet:", sheet)
        break

if df is None:
    raise ValueError(
        "No Sales column found. Check the Excel column names."
    )

# Normalize column names
df.columns = [str(c).strip() for c in df.columns]

def find_column(possible_names):
    for col in df.columns:
        if col.strip().lower() in possible_names:
            return col
    return None

sales_col = find_column({"sales", "revenue", "total sales"})
profit_col = find_column({"profit", "net profit"})
date_col = find_column({"order date", "date", "sales date"})
category_col = find_column({"category", "product category"})
region_col = find_column({"region", "sales region"})
product_col = find_column(
    {"product name", "product", "item", "item name"}
)
quantity_col = find_column({"quantity", "qty", "units sold"})
discount_col = find_column({"discount", "discount rate"})
order_col = find_column({"order id", "order number", "order no"})

# Convert numeric columns safely
df[sales_col] = pd.to_numeric(df[sales_col], errors="coerce")
df = df.dropna(subset=[sales_col]).copy()

if profit_col:
    df[profit_col] = pd.to_numeric(df[profit_col], errors="coerce")
else:
    df["Profit"] = np.nan
    profit_col = "Profit"

if quantity_col:
    df[quantity_col] = pd.to_numeric(
        df[quantity_col], errors="coerce"
    )
else:
    df["Quantity"] = 1
    quantity_col = "Quantity"

if discount_col:
    df[discount_col] = pd.to_numeric(
        df[discount_col], errors="coerce"
    )
else:
    df["Discount"] = 0.0
    discount_col = "Discount"

if date_col:
    df[date_col] = pd.to_datetime(df[date_col], errors="coerce")
else:
    df["Order Date"] = pd.NaT
    date_col = "Order Date"

if category_col is None:
    df["Category"] = "Unknown"
    category_col = "Category"

if region_col is None:
    df["Region"] = "Unknown"
    region_col = "Region"

if product_col is None:
    df["Product"] = "Unknown"
    product_col = "Product"

if order_col is None:
    df["Order ID"] = np.arange(1, len(df) + 1)
    order_col = "Order ID"


total_sales = float(df[sales_col].sum())
total_profit = float(df[profit_col].sum()) if df[profit_col].notna().any() else 0
total_orders = int(df[order_col].nunique())
total_quantity = float(df[quantity_col].sum())
profit_margin = (total_profit / total_sales * 100) if total_sales else 0

# Category analysis
category_analysis = (
    df.groupby(category_col, dropna=False)
    .agg(
        Total_Sales=(sales_col, "sum"),
        Total_Profit=(profit_col, "sum"),
        Total_Quantity=(quantity_col, "sum"),
        Records=(sales_col, "size")
    )
    .reset_index()
    .sort_values("Total_Sales", ascending=False)
)


region_analysis = (
    df.groupby(region_col, dropna=False)
    .agg(
        Total_Sales=(sales_col, "sum"),
        Total_Profit=(profit_col, "sum"),
        Records=(sales_col, "size")
    )
    .reset_index()
    .sort_values("Total_Sales", ascending=False)
)


product_analysis = (
    df.groupby(product_col, dropna=False)
    .agg(
        Total_Sales=(sales_col, "sum"),
        Total_Profit=(profit_col, "sum"),
        Total_Quantity=(quantity_col, "sum")
    )
    .reset_index()
    .sort_values("Total_Sales", ascending=False)
)


dated = df.dropna(subset=[date_col]).copy()

if not dated.empty:
    dated["Month"] = dated[date_col].dt.to_period("M").astype(str)
    monthly_analysis = (
        dated.groupby("Month")
        .agg(
            Total_Sales=(sales_col, "sum"),
            Total_Profit=(profit_col, "sum")
        )
        .reset_index()
        .sort_values("Month")
    )
else:
    monthly_analysis = pd.DataFrame(
        columns=["Month", "Total_Sales", "Total_Profit"]
    )


numeric_df = df.select_dtypes(include=np.number)
statistics_summary = numeric_df.describe().T.reset_index()
statistics_summary = statistics_summary.rename(
    columns={"index": "Column"}
)

missing_values = df.isna().sum().rename("Missing Values").reset_index()
missing_values.columns = ["Column", "Missing Values"]


correlation = numeric_df.corr()


if (
    sales_col in numeric_df.columns
    and profit_col in numeric_df.columns
    and numeric_df[sales_col].nunique() > 1
    and numeric_df[profit_col].nunique() > 1
    and len(numeric_df) >= 3
):
    valid = df[[sales_col, profit_col]].dropna()
    if len(valid) >= 3:
        pearson_r, pearson_p = stats.pearsonr(
            valid[sales_col], valid[profit_col]
        )
    else:
        pearson_r, pearson_p = np.nan, np.nan
else:
    pearson_r, pearson_p = np.nan, np.nan


sales_symbol, profit_symbol = sp.symbols("Sales Profit")
margin_formula = sp.simplify(
    profit_symbol / sales_symbol * 100
)

formula_analysis = pd.DataFrame({
    "Item": [
        "Profit Margin Formula",
        "Calculated Total Sales",
        "Calculated Total Profit",
        "Profit Margin (%)"
    ],
    "Value": [
        str(margin_formula),
        total_sales,
        total_profit,
        profit_margin
    ]
})

# ---------- 3. MACHINE LEARNING ----------
ml_summary = pd.DataFrame({
    "Metric": ["Status", "R2 Score", "RMSE"],
    "Value": ["Not enough data", np.nan, np.nan]
})

features = [sales_col, quantity_col, discount_col]

ml_data = df[features + [profit_col]].dropna().copy()

if (
    profit_col != sales_col
    and len(ml_data) >= 10
    and ml_data[profit_col].nunique() > 1
):
    X = ml_data[features]
    y = ml_data[profit_col]

    X_train, X_test, y_train, y_test = train_test_split(
        X, y, test_size=0.2, random_state=42
    )

    model = LinearRegression()
    model.fit(X_train, y_train)
    predicted = model.predict(X_test)

    ml_summary = pd.DataFrame({
        "Metric": ["Status", "R2 Score", "RMSE"],
        "Value": [
            "Model trained",
            r2_score(y_test, predicted),
            np.sqrt(mean_squared_error(y_test, predicted))
        ]
    })

# ---------- 4. CREATE CHARTS ----------
chart_paths = {}

# Sales by category
fig, ax = plt.subplots(figsize=(8, 4.5))
sns.barplot(
    data=category_analysis,
    x=category_col,
    y="Total_Sales",
    ax=ax
)
ax.set_title("Sales by Category")
ax.set_xlabel("Category")
ax.set_ylabel("Total Sales")
plt.xticks(rotation=20)
fig.tight_layout()
chart_paths["category"] = CHART_DIR / "sales_by_category.png"
fig.savefig(chart_paths["category"], dpi=150)
plt.close(fig)

# Sales by region
fig, ax = plt.subplots(figsize=(8, 4.5))
sns.barplot(
    data=region_analysis,
    x=region_col,
    y="Total_Sales",
    ax=ax
)
ax.set_title("Sales by Region")
ax.set_xlabel("Region")
ax.set_ylabel("Total Sales")
plt.xticks(rotation=20)
fig.tight_layout()
chart_paths["region"] = CHART_DIR / "sales_by_region.png"
fig.savefig(chart_paths["region"], dpi=150)
plt.close(fig)

# Monthly sales trend
if not monthly_analysis.empty:
    fig, ax = plt.subplots(figsize=(9, 4.5))
    ax.plot(
        monthly_analysis["Month"],
        monthly_analysis["Total_Sales"],
        marker="o",
        label="Sales"
    )
    ax.set_title("Monthly Sales Trend")
    ax.set_xlabel("Month")
    ax.set_ylabel("Sales")
    ax.tick_params(axis="x", rotation=45)
    ax.legend()
    fig.tight_layout()
    chart_paths["monthly"] = CHART_DIR / "monthly_sales.png"
    fig.savefig(chart_paths["monthly"], dpi=150)
    plt.close(fig)

# Correlation heatmap
if not correlation.empty:
    fig, ax = plt.subplots(figsize=(8, 5))
    sns.heatmap(
        correlation,
        annot=True,
        fmt=".2f",
        cmap="coolwarm",
        ax=ax
    )
    ax.set_title("Numeric Data Correlation")
    fig.tight_layout()
    chart_paths["heatmap"] = CHART_DIR / "correlation_heatmap.png"
    fig.savefig(chart_paths["heatmap"], dpi=150)
    plt.close(fig)

# ---------- 5. EXPORT DATA TO EXCEL ----------
dashboard_summary = pd.DataFrame({
    "KPI": [
        "Total Sales",
        "Total Profit",
        "Unique Orders",
        "Total Quantity",
        "Profit Margin (%)",
        "Number of Records"
    ],
    "Value": [
        total_sales,
        total_profit,
        total_orders,
        total_quantity,
        profit_margin,
        len(df)
    ]
})

with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
    dashboard_summary.to_excel(
        writer, sheet_name="Dashboard", index=False
    )
    df.to_excel(writer, sheet_name="Sales Data", index=False)
    category_analysis.to_excel(
        writer, sheet_name="Category Analysis", index=False
    )
    region_analysis.to_excel(
        writer, sheet_name="Region Analysis", index=False
    )
    monthly_analysis.to_excel(
        writer, sheet_name="Monthly Analysis", index=False
    )
    statistics_summary.to_excel(
        writer, sheet_name="Statistics", index=False
    )
    missing_values.to_excel(
        writer, sheet_name="Missing Values", index=False
    )
    correlation.to_excel(writer, sheet_name="Correlation")
    formula_analysis.to_excel(
        writer, sheet_name="Formula Analysis", index=False
    )
    ml_summary.to_excel(
        writer, sheet_name="ML Results", index=False
    )

# ---------- 6. FORMAT EXCEL DASHBOARD ----------
wb = load_workbook(OUTPUT_FILE)
ws = wb["Dashboard"]

ws["D1"] = "SALES ANALYSIS DASHBOARD"
ws["D1"].font = Font(size=18, bold=True, color="FFFFFF")
ws["D1"].fill = PatternFill(
    "solid", fgColor="1F4E78"
)
ws["D1"].alignment = Alignment(horizontal="center")

ws.merge_cells("D1:H1")

# Format summary table
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=2):
    row[0].font = Font(bold=True)
    row[1].number_format = '#,##0.00'

ws.column_dimensions["A"].width = 25
ws.column_dimensions["B"].width = 20
ws.column_dimensions["D"].width = 18

# Insert charts in Dashboard sheet
positions = {
    "category": "A10",
    "region": "J10",
    "monthly": "A30",
    "heatmap": "J30"
}

for key, image_path in chart_paths.items():
    img = XLImage(str(image_path))
    img.width = 560
    img.height = 315
    ws.add_image(img, positions[key])

# Improve readability of data sheets
for sheet in wb.worksheets:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")

    for col_cells in sheet.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_length = 0

        for cell in col_cells[:100]:
            value = str(cell.value) if cell.value is not None else ""
            max_length = max(max_length, len(value))

        sheet.column_dimensions[col_letter].width = min(
            max(max_length + 2, 12), 35
        )

# Restore dashboard chart placement after column sizing
ws.column_dimensions["A"].width = 25
ws.column_dimensions["B"].width = 20
ws.column_dimensions["D"].width = 18

wb.save(OUTPUT_FILE)

print("\nSALES ANALYSIS COMPLETED!")
print("Total records:", len(df))
print("Total sales:", round(total_sales, 2))
print("Total profit:", round(total_profit, 2))
print("Unique orders:", total_orders)
print("Excel dashboard saved to:")
print(OUTPUT_FILE)
